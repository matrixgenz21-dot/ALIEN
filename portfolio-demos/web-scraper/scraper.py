"""
Web Scraper Tool - Kisi bhi website ka data Excel mein save karo
By: Matrix (Fiverr Portfolio Demo)

Features:
- Any website se data scrape karo
- Tables, lists, headings, links extract karo
- Excel file mein save karo
- CSV export bhi available
- Command line se chalao — easy to use

Usage:
    python scraper.py https://example.com
    python scraper.py https://example.com --output data.xlsx
    python scraper.py https://example.com --format csv
"""

import sys
import os
import re
import json
from datetime import datetime

try:
    import requests
    REQUESTS_AVAILABLE = True
except ImportError:
    REQUESTS_AVAILABLE = False

try:
    from bs4 import BeautifulSoup
    BS4_AVAILABLE = True
except ImportError:
    BS4_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class WebScraper:
    """Professional web scraper — any website ka data extract karo."""

    def __init__(self):
        self.headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        }
        self.data = {}

    def scrape(self, url):
        """Website se sab data extract karo."""
        if not REQUESTS_AVAILABLE:
            print("[ERROR] requests not installed: pip install requests")
            return None
        if not BS4_AVAILABLE:
            print("[ERROR] beautifulsoup4 not installed: pip install beautifulsoup4")
            return None

        print(f"\n[Scraper] Fetching: {url}")
        try:
            response = requests.get(url, headers=self.headers, timeout=15)
            response.raise_for_status()
        except requests.RequestException as e:
            print(f"[ERROR] Could not fetch URL: {e}")
            return None

        soup = BeautifulSoup(response.text, "html.parser")
        print(f"[Scraper] Page loaded. Extracting data...")

        self.data = {
            "url": url,
            "title": soup.title.string.strip() if soup.title and soup.title.string else "No Title",
            "scraped_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "headings": self._extract_headings(soup),
            "links": self._extract_links(soup, url),
            "images": self._extract_images(soup, url),
            "tables": self._extract_tables(soup),
            "paragraphs": self._extract_paragraphs(soup),
            "meta": self._extract_meta(soup),
        }

        # Summary
        print(f"\n--- Scrape Summary ---")
        print(f"  Title: {self.data['title']}")
        print(f"  Headings: {len(self.data['headings'])}")
        print(f"  Links: {len(self.data['links'])}")
        print(f"  Images: {len(self.data['images'])}")
        print(f"  Tables: {len(self.data['tables'])}")
        print(f"  Paragraphs: {len(self.data['paragraphs'])}")
        print(f"----------------------\n")

        return self.data

    def _extract_headings(self, soup):
        headings = []
        for tag in ["h1", "h2", "h3", "h4", "h5", "h6"]:
            for h in soup.find_all(tag):
                text = h.get_text(strip=True)
                if text:
                    headings.append({"level": tag, "text": text})
        return headings

    def _extract_links(self, soup, base_url):
        links = []
        for a in soup.find_all("a", href=True):
            href = a["href"]
            text = a.get_text(strip=True) or "No Text"
            if href.startswith("/"):
                href = base_url.rstrip("/") + href
            if href.startswith("http"):
                links.append({"text": text[:100], "url": href})
        return links

    def _extract_images(self, soup, base_url):
        images = []
        for img in soup.find_all("img", src=True):
            src = img["src"]
            alt = img.get("alt", "No Alt Text")
            if src.startswith("/"):
                src = base_url.rstrip("/") + src
            images.append({"alt": alt[:100], "src": src})
        return images

    def _extract_tables(self, soup):
        tables = []
        for table in soup.find_all("table"):
            rows = []
            for tr in table.find_all("tr"):
                cells = []
                for td in tr.find_all(["td", "th"]):
                    cells.append(td.get_text(strip=True))
                if cells:
                    rows.append(cells)
            if rows:
                tables.append(rows)
        return tables

    def _extract_paragraphs(self, soup):
        paragraphs = []
        for p in soup.find_all("p"):
            text = p.get_text(strip=True)
            if text and len(text) > 20:
                paragraphs.append(text[:500])
        return paragraphs

    def _extract_meta(self, soup):
        meta = {}
        for tag in soup.find_all("meta"):
            name = tag.get("name", tag.get("property", ""))
            content = tag.get("content", "")
            if name and content:
                meta[name] = content[:200]
        return meta

    def save_excel(self, filename="scraped_data.xlsx"):
        """Data ko Excel file mein save karo."""
        if not OPENPYXL_AVAILABLE:
            print("[ERROR] openpyxl not installed: pip install openpyxl")
            return False

        if not self.data:
            print("[ERROR] No data to save. Run scrape() first.")
            return False

        wb = openpyxl.Workbook()

        # Sheet 1: Summary
        ws = wb.active
        ws.title = "Summary"
        ws.append(["Property", "Value"])
        ws.append(["URL", self.data["url"]])
        ws.append(["Title", self.data["title"]])
        ws.append(["Scraped At", self.data["scraped_at"]])
        ws.append(["Total Headings", len(self.data["headings"])])
        ws.append(["Total Links", len(self.data["links"])])
        ws.append(["Total Images", len(self.data["images"])])
        ws.append(["Total Tables", len(self.data["tables"])])

        # Style header
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)

        # Sheet 2: Headings
        ws2 = wb.create_sheet("Headings")
        ws2.append(["Level", "Text"])
        for h in self.data["headings"]:
            ws2.append([h["level"], h["text"]])
        for cell in ws2[1]:
            cell.font = openpyxl.styles.Font(bold=True)

        # Sheet 3: Links
        ws3 = wb.create_sheet("Links")
        ws3.append(["Text", "URL"])
        for link in self.data["links"][:500]:
            ws3.append([link["text"], link["url"]])
        for cell in ws3[1]:
            cell.font = openpyxl.styles.Font(bold=True)

        # Sheet 4: Images
        ws4 = wb.create_sheet("Images")
        ws4.append(["Alt Text", "Source URL"])
        for img in self.data["images"][:500]:
            ws4.append([img["alt"], img["src"]])
        for cell in ws4[1]:
            cell.font = openpyxl.styles.Font(bold=True)

        # Sheet 5: Tables
        for i, table in enumerate(self.data["tables"][:10]):
            ws_t = wb.create_sheet(f"Table_{i+1}")
            for row in table:
                ws_t.append(row)
            if ws_t[1]:
                for cell in ws_t[1]:
                    cell.font = openpyxl.styles.Font(bold=True)

        # Sheet 6: Content
        ws5 = wb.create_sheet("Content")
        ws5.append(["Paragraph"])
        for p in self.data["paragraphs"][:200]:
            ws5.append([p])
        for cell in ws5[1]:
            cell.font = openpyxl.styles.Font(bold=True)

        # Auto-width columns
        for ws_sheet in wb.worksheets:
            for col in ws_sheet.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except (TypeError, AttributeError):
                        pass
                ws_sheet.column_dimensions[col_letter].width = min(max_length + 2, 60)

        wb.save(filename)
        print(f"[Scraper] Data saved to: {filename}")
        return True

    def save_csv(self, filename="scraped_data.csv"):
        """Data ko CSV file mein save karo."""
        import csv
        if not self.data:
            print("[ERROR] No data to save.")
            return False

        with open(filename, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Type", "Content", "Extra"])

            for h in self.data["headings"]:
                writer.writerow(["Heading", h["text"], h["level"]])
            for link in self.data["links"]:
                writer.writerow(["Link", link["text"], link["url"]])
            for img in self.data["images"]:
                writer.writerow(["Image", img["alt"], img["src"]])
            for p in self.data["paragraphs"]:
                writer.writerow(["Paragraph", p, ""])

        print(f"[Scraper] Data saved to: {filename}")
        return True

    def save_json(self, filename="scraped_data.json"):
        """Data ko JSON mein save karo."""
        if not self.data:
            return False
        with open(filename, "w", encoding="utf-8") as f:
            json.dump(self.data, f, indent=2, ensure_ascii=False)
        print(f"[Scraper] Data saved to: {filename}")
        return True


def main():
    if len(sys.argv) < 2:
        print("=" * 50)
        print("  Web Scraper Tool")
        print("  By: Matrix")
        print("=" * 50)
        print("\nUsage:")
        print("  python scraper.py <URL>")
        print("  python scraper.py <URL> --output filename.xlsx")
        print("  python scraper.py <URL> --format csv")
        print("\nExamples:")
        print("  python scraper.py https://example.com")
        print("  python scraper.py https://news.ycombinator.com --output hn.xlsx")
        return

    url = sys.argv[1]
    output = "scraped_data"
    fmt = "xlsx"

    for i, arg in enumerate(sys.argv[2:], 2):
        if arg == "--output" and i + 1 < len(sys.argv):
            output = sys.argv[i + 1].rsplit(".", 1)[0]
        if arg == "--format" and i + 1 < len(sys.argv):
            fmt = sys.argv[i + 1].lower()

    scraper = WebScraper()
    data = scraper.scrape(url)

    if data:
        if fmt == "csv":
            scraper.save_csv(f"{output}.csv")
        elif fmt == "json":
            scraper.save_json(f"{output}.json")
        else:
            scraper.save_excel(f"{output}.xlsx")
        print("\nDone! Data extracted and saved successfully.")


if __name__ == "__main__":
    main()
